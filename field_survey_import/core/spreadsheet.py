"""Spreadsheet + photos bundle export (scripts/export_spreadsheet.py's engine): a
Field Survey zip -> one zip holding an .xlsx and a .csv of a curated subset of the
observation properties, every referenced photo (original bytes, untouched), and any
extra files the caller wants shipped alongside (typically the HTML map from
scripts/export_html.py).

Zero qgis/PyQt5 imports. openpyxl (for the .xlsx) is imported lazily inside
`write_xlsx`, the same pattern and reason as Pillow in core/photo_optimize.py, so
this module stays importable without it. Timezone conversion uses the stdlib
`zoneinfo` - which on Windows has no system tz database to fall back on, so the
`tzdata` package must be installed (requirements.txt) or `ZoneInfo("Europe/London")`
raises `ZoneInfoNotFoundError`.
"""
from __future__ import annotations  # `X | None` unions must stay lazy on Python 3.9

import csv
import datetime as dt
import tempfile
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from .media import MediaRef, resolve_media
from .model import Observation, SurveyExport

DEFAULT_TIMEZONE = "Europe/London"

# The curated subset the spreadsheet carries, in column order. Deliberately NOT the
# full 26-field schema (core/schema.py) - this export is for people reading a
# spreadsheet, not for round-tripping data.
COLUMNS: tuple[str, ...] = (
    "note", "lat", "lon", "os_grid_ref", "photo", "heading_deg", "direction",
    "recorded_date", "recorded_time",
)

# 16-point compass, clockwise from north, each sector 22.5 degrees wide and centred
# on its named bearing (so "n" covers 348.75 up to 11.25).
COMPASS_16: tuple[str, ...] = (
    "n", "nne", "ne", "ene", "e", "ese", "se", "sse",
    "s", "ssw", "sw", "wsw", "w", "wnw", "nw", "nnw",
)
_SECTOR_DEG = 360 / len(COMPASS_16)

_PHOTO_SEPARATOR = "; "  # multi-photo observations list every filename in one cell


def compass_direction(heading_deg: float | None) -> str:
    """Nearest 16-point compass name for a bearing (degrees clockwise from north),
    lowercase; '' for a null heading.
    """
    if heading_deg is None:
        return ""
    # floor(x + half) rather than round(): round() is banker's rounding, which
    # would send an exact sector boundary like 11.25 down to "n" instead of "nne".
    index = int((heading_deg % 360 + _SECTOR_DEG / 2) // _SECTOR_DEG) % len(COMPASS_16)
    return COMPASS_16[index]


def _blank_if_none(value: Any) -> Any:
    return "" if value is None else value


def _photo_cell(obs: Observation) -> str:
    return _PHOTO_SEPARATOR.join(entry.photo for entry in obs.photos)


def build_rows(export: SurveyExport, *, tz: str = DEFAULT_TIMEZONE) -> list[dict[str, Any]]:
    """One dict per observation, keyed by COLUMNS in order, values kept typed
    (floats, `datetime.date`, `datetime.time`) so the xlsx writer can emit real
    cells; `csv_row()` stringifies them for the .csv. `recorded_at` is converted from
    its UTC instant into `tz` (an IANA name) before being split into date and time.
    """
    zone = ZoneInfo(tz)
    rows: list[dict[str, Any]] = []
    for obs in export.observations:
        local = obs.recorded_at.astimezone(zone)
        rows.append({
            "note": obs.note,
            "lat": obs.lat,
            "lon": obs.lon,
            "os_grid_ref": _blank_if_none(obs.os_grid_ref),
            "photo": _photo_cell(obs),
            "heading_deg": _blank_if_none(obs.heading_deg),
            "direction": compass_direction(obs.heading_deg),
            "recorded_date": local.date(),
            "recorded_time": local.time().replace(microsecond=0),
        })
    return rows


def csv_row(row: dict[str, Any]) -> dict[str, str]:
    """Stringify one build_rows() row for csv.DictWriter: ISO date, HH:MM:SS time,
    repr-precision floats (so lat/lon survive a round trip), everything else str().
    """
    out: dict[str, str] = {}
    for key, value in row.items():
        if isinstance(value, dt.date):  # dt.datetime is a dt.date subclass too
            out[key] = value.isoformat()
        elif isinstance(value, dt.time):
            out[key] = value.strftime("%H:%M:%S")
        elif isinstance(value, float):
            out[key] = repr(value)
        else:
            out[key] = str(value)
    return out


def write_csv(rows: Iterable[dict[str, Any]], path: Path) -> None:
    """utf-8-sig (BOM) so Excel on Windows opens non-ASCII notes as UTF-8 rather than
    guessing the system code page; newline='' per the csv module's own docs.
    """
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow(csv_row(row))


_XLSX_NUMBER_FORMATS = {"recorded_date": "yyyy-mm-dd", "recorded_time": "hh:mm:ss"}
_XLSX_MIN_WIDTH = 10
_XLSX_MAX_WIDTH = 60


def write_xlsx(rows: Iterable[dict[str, Any]], path: Path) -> None:
    """Single sheet, bold header, date/time cells typed (not text) with explicit
    number formats, column widths sized to content. openpyxl is imported here, not
    at module level - see the module docstring.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "observations"
    ws.append(list(COLUMNS))
    for cell in ws[1]:
        cell.font = Font(bold=True)

    widths = {col: len(col) for col in COLUMNS}
    for row in rows:
        ws.append([row[col] for col in COLUMNS])
        for col in COLUMNS:
            widths[col] = max(widths[col], len(csv_row(row)[col]))
    for col, fmt in _XLSX_NUMBER_FORMATS.items():
        col_index = COLUMNS.index(col) + 1
        for cell in ws.iter_cols(min_col=col_index, max_col=col_index, min_row=2):
            for c in cell:
                c.number_format = fmt
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(
            _XLSX_MAX_WIDTH, max(_XLSX_MIN_WIDTH, widths[col] + 2)
        )
    ws.freeze_panes = "A2"
    wb.save(path)


@dataclass(frozen=True)
class BundleSummary:
    session_name: str
    row_count: int
    photo_count: int
    included_files: tuple[str, ...]
    output_size_bytes: int


def _photo_refs(export: SurveyExport, zf: zipfile.ZipFile) -> list[MediaRef]:
    """Every photo across the export, deduped by zip entry, in first-seen order.
    Raises MediaJoinError (via resolve_media) if any is missing from the zip.
    """
    seen: dict[str, MediaRef] = {}
    for obs in export.observations:
        photo_refs, _audio = resolve_media(obs, zf)
        for ref in photo_refs:
            seen.setdefault(ref.zip_entry, ref)
    return list(seen.values())


def write_bundle(
    export: SurveyExport,
    zf: zipfile.ZipFile,
    out_path: Path,
    *,
    name: str,
    tz: str = DEFAULT_TIMEZONE,
    include: Iterable[Path] = (),
) -> BundleSummary:
    """Write `<name>.xlsx`, `<name>.csv`, every photo under `photos/`, and each
    `include` file at its basename into a single zip at out_path. Photo bytes are
    copied straight out of the source zip (zf.read), never optimised - this bundle
    is the "here are the actual photos" deliverable, unlike the KMZ/HTML exports.
    Member names are plain forward-slash strings, never Path joins (backslashes on
    Windows silently corrupt zip member names - see CLAUDE.md).
    """
    include = [Path(p) for p in include]
    spreadsheet_names = {f"{name}.xlsx", f"{name}.csv"}
    included_names: list[str] = []
    for path in include:
        if path.name in spreadsheet_names or path.name in included_names:
            raise ValueError(
                f"--include {path} would overwrite {path.name!r} inside the bundle"
            )
        included_names.append(path.name)

    rows = build_rows(export, tz=tz)
    photo_refs = _photo_refs(export, zf)  # fail on a missing photo before writing anything

    with tempfile.TemporaryDirectory() as tmp:
        xlsx_path = Path(tmp) / f"{name}.xlsx"
        csv_path = Path(tmp) / f"{name}.csv"
        write_xlsx(rows, xlsx_path)
        write_csv(rows, csv_path)
        with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as out:
            out.write(xlsx_path, xlsx_path.name)
            out.write(csv_path, csv_path.name)
            for path in include:
                out.write(path, path.name)
            for ref in photo_refs:
                out.writestr(f"photos/{ref.filename}", zf.read(ref.zip_entry))

    return BundleSummary(
        session_name=export.session.name,
        row_count=len(rows),
        photo_count=len(photo_refs),
        included_files=tuple(included_names),
        output_size_bytes=out_path.stat().st_size,
    )
